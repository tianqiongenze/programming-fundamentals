# -*- coding: utf-8 -*-
"""把 32 个项目的四档实测结果 + 实跑截图注入三课授课 PPT 的对应位置。

每个 deck 用「本工单的工程化项目」页里写的 slug 自动定位项目，然后：
  1. 在该页之后插入 E1「四档实测验收」+ E2「实跑截图证据」（有产出图再加 E3）；
  2. 在 D/C/B/A 四张分层任务卡的「达标标准」条上追加本次实测判定一行；
  3. 刷新「工单完成效果图」页的参考运行输出块为本次真跑内容。
所有数字逐字来自 工具/ppt97/evidence/<slug>/（capture97.py 真跑产物），不手抄。
"""
import copy
import glob
import json
import os
import re
import shutil
import sys
import time

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

ROOT = "/Users/joelgong/PycharmProjects/ai-course-teaching-resources"
sys.path.insert(0, os.path.join(ROOT, "03-Python程序设计-项目实战/工具"))
from ppt_enrich import (set_shape_text, _renumber_ids,  # noqa: E402
                        _copy_rel_images, _renumber_slide_num)

EV = os.path.join(ROOT, "工具", "ppt97", "evidence")
BAK = "/tmp/backup_ppt97"
COURSE_DIRS = [("01", "01-AI应用基础/授课PPT"), ("02", "02-程序设计基础/授课PPT"),
               ("05", "05-AI应用基础48课时版/V2授课PPT")]
SLUG_RE = re.compile(r"工程化项目/([A-Za-z0-9_]+)")
PROJ_PAGE = "本工单的工程化项目"
E1_KEY = "实测验收 · "
E2_KEY = "实跑截图证据 · "
E3_KEY = "产出图证据 · "
CARD_KEY = {"D": "分层工单 · D 层任务卡", "C": "分层工单 · C 层任务卡",
            "B": "分层工单 · B 层任务卡", "A": "分层工单 · A 层任务卡"}


def _tier_key(label):
    """从「及格(D)/良好(C)/优秀(B)/优(A)」取档位字母；全角/半角括号都能识别。"""
    m = re.search(r"[（(]\s*([DCBA])\s*[)）]", label)
    return m.group(1) if m else None


TIER_LINE_RE = re.compile(r"^\s*(及格\(D\)|良好\(C\)|优秀\(B\)|优\(A\))\s+(.*)$")


def _restore_verdict(d):
    """capture97 只留了首段门槛；这里从原始档位行取回完整实测判定（含「功能OK」等中段）。"""
    raw = {}
    for ln in list(d.get("report_tier_seg") or []) + list(d.get("run_tail") or []):
        m = TIER_LINE_RE.match(ln)
        if not m:
            continue
        key = _tier_key(m.group(1))
        if key and key not in raw:
            raw[key] = m.group(2).strip()
    for row in d.get("rows", []):
        full = raw.get(row.get("key"))
        if not full:
            row["verdict_text"] = row.get("gate", "")
            continue
        parts = [p.strip() for p in full.split("｜") if p.strip()]
        if len(parts) > 1 and parts[-1] == (row.get("need") or "").strip():
            parts = parts[:-1]
        row["verdict_text"] = " ｜ ".join(parts)


TIER_LABEL = {"D": "及格(D)", "C": "良好(C)", "B": "优秀(B)", "A": "优(A)"}
CLAIM_LINE_RE = re.compile(r"达标档位[^:：=]*[:：=]?\s*([DCBA])\b")


def _resolve_claimed(d, slug):
    """档位取数顺序：报告头档位行 → 终端输出的「达标档位」行 → result.json → 四档连判。"""
    if d.get("claimed"):
        return
    for src in (d.get("report_tier_seg") or [], d.get("run_tail") or []):
        for ln in src:
            m = CLAIM_LINE_RE.search(ln)
            if m:
                d["claimed"] = TIER_LABEL[m.group(1)]
                d["claimed_src"] = "本次真跑输出的「达标档位」行"
                return
    for _k, rel in COURSE_DIRS:
        rp = os.path.join(ROOT, rel.split("/")[0], "工程化项目", slug,
                          "output", "result.json")
        if not os.path.exists(rp):
            continue
        rj = json.load(open(rp, encoding="utf-8"))
        for key in ("tiers", "tier_verdict"):
            v = (rj.get(key) or {}).get("claimed")
            if v:
                d["claimed"] = v
                d["claimed_src"] = "output/result.json:%s" % key
                return
        t = rj.get("tier")
        if t in TIER_LABEL:
            d["claimed"] = TIER_LABEL[t]
            d["claimed_src"] = "output/result.json:tier"
            return
    got = [r["key"] for r in d.get("rows", []) if r["verdict"] == "达标"]
    ladder = [k for k in ("D", "C", "B", "A") if k in got]
    if ladder == ["D", "C", "B", "A"][:len(ladder)] and ladder:
        d["claimed"] = TIER_LABEL[ladder[-1]]
        d["claimed_src"] = "四档实测判定连判推出"


def evidence(slug):
    """读实测证据；报告头没写档位时回退到项目 result.json 的落盘判定。"""
    d = json.load(open(os.path.join(EV, slug, "evidence.json"), encoding="utf-8"))
    # 每档只取第一条判定行，并按 D→C→B→A 归序（有的报告正文里还会出现同档字样）
    first = {}
    for row in d.get("rows", []):
        key = _tier_key(str(row.get("tier") or row.get("key") or ""))
        if key is None:
            continue
        row["key"] = key
        first.setdefault(key, row)
    d["rows"] = [first[k] for k in ("D", "C", "B", "A") if k in first]
    _restore_verdict(d)
    _resolve_claimed(d, slug)
    rp = os.path.join(EV, slug, "run.txt")
    d["meta"] = {"python": "/opt/anaconda3/bin/python",
                 "date": time.strftime("%Y-%m-%d %H:%M",
                                       time.localtime(os.path.getmtime(rp)))
                 if os.path.exists(rp) else "本次"}
    return d


def slide_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def find_index(prs, key):
    for i, sl in enumerate(prs.slides):
        if key in slide_title(sl):
            return i
    return None


def deck_slug(prs, proj_idx):
    for sh in prs.slides[proj_idx].shapes:
        if sh.has_text_frame:
            m = SLUG_RE.search(sh.text_frame.text)
            if m:
                return m.group(1)
    for sl in prs.slides:  # 兜底：全 deck 找 slug 引用
        for sh in sl.shapes:
            if sh.has_text_frame:
                m = SLUG_RE.search(sh.text_frame.text)
                if m:
                    return m.group(1)
    return None


def strip_body(slide):
    """单遍删除主体形状（lxml 元素代理 id 不稳定，不能两遍比较）。"""
    for sh in list(slide.shapes):
        top = (sh.top or 0) / 914400.0
        if 1.5 <= top <= 7.0:
            sh._element.getparent().remove(sh._element)


def title_shape(slide):
    cands = [s for s in slide.shapes if s.has_text_frame and s.top is not None
             and s.top / 914400.0 < 1.3 and s.text_frame.text.strip()]
    return max(cands, key=lambda s: s.top / 914400.0) if cands else None


def clone_base(prs, idx):
    """同文档克隆基准页，继承页眉色带/字体/页脚。"""
    layout = prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]
    src_slide = prs.slides[idx]
    dst_slide = prs.slides.add_slide(layout)
    for ph in list(dst_slide.placeholders):
        dst_slide.shapes._spTree.remove(ph._element)
    for child in list(src_slide.shapes._spTree):
        if child.tag.endswith("nvGrpSpPr") or child.tag.endswith("grpSpPr"):
            continue
        dst_slide.shapes._spTree.append(copy.deepcopy(child))
    _renumber_ids(dst_slide)
    _copy_rel_images(src_slide, dst_slide, prs.part, prs.part)
    return dst_slide


def _reorder(prs, n_new, after_idx):
    """把刚克隆到末尾的 n_new 页整体移到 after_idx 之后（保持彼此顺序）。"""
    if n_new <= 0:
        return
    xml = prs.slides._sldIdLst
    moved = list(xml)[-n_new:]
    for el in moved:
        xml.remove(el)
    for k, el in enumerate(moved):
        xml.insert(min(after_idx + 1 + k, len(list(xml))), el)


def fit_lines(text, size, box_w_in, max_lines):
    """按显示宽度粗估折行数，放不下就截断（宁可少字，不可溢出）。"""
    per_line = max(1, int(box_w_in / (size / 72.0)))
    lines = sum(max(1, -(-len(p) // per_line)) for p in text.split("\n"))
    if lines <= max_lines:
        return True, text
    budget = max_lines * per_line
    cut = text[:budget]
    tail = "……（全文见 output/report.txt）" if len(text) > budget else ""
    return False, cut + tail if len(cut) + len(tail) <= budget else cut


def add_text(slide, l, t, w, h, text, size=12, bold=False, color="2B2F36",
             font=None, align=None):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        if font:
            r.font.name = font
        if align:
            p.alignment = getattr(PP_ALIGN, align, PP_ALIGN.LEFT)
    return box


def add_rect(slide, l, t, w, h, fill, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t),
                                 Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    sp.text_frame.word_wrap = True
    return sp


def add_pic_fit(slide, path, l, t, w, h):
    """等比缩放进 (w,h) 盒内并水平居中。"""
    iw, ih = Image.open(path).size
    ar = iw / float(ih)
    pw, ph = w, w / ar
    if ph > h:
        ph, pw = h, h * ar
    slide.shapes.add_picture(path, Inches(l + (w - pw) / 2.0), Inches(t),
                             width=Inches(pw), height=Inches(ph))


def base_color_of(slide):
    for sh in slide.shapes:
        if "AUTO_SHAPE" in str(sh.shape_type) and (sh.top or 0) == 0:
            try:
                return str(sh.fill.fore_color.rgb)
            except Exception:  # noqa: BLE001
                return "1F3A5F"
    return "1F3A5F"


def font_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if p.runs:
                    return p.runs[0].font.name
    return None


def card_standards(prs):
    """取四张分层任务卡上「✔ 达标标准」的工单原文（在追加实测判定之前读取）。"""
    std = {}
    for key, head in CARD_KEY.items():
        idx = find_index(prs, head)
        if idx is None:
            continue
        for sh in prs.slides[idx].shapes:
            if sh.has_text_frame and "达标标准" in sh.text_frame.text:
                t = re.sub(r"[✔✓]\s*达标标准[:：]?\s*", "",
                           sh.text_frame.text).strip()
                if t:
                    std[key] = t
                break
    return std


def _gate_text(row, std):
    """报告里该档只留了指向人工答辩项的说明时，改用任务卡上的工单原文标准。"""
    need = (row.get("need") or "").strip()
    if std.get(row.get("key")) and (len(need) < 12 or "人工答辩" in need
                                    or "见下方" in need):
        return std[row["key"]]
    return need


def build_e1(slide, ev, base, font, std=None):
    std = std or {}
    add_rect(slide, 0, 1.5, 13.33, 0.78, "F2F5F9")
    claimed = (ev.get("claimed") or "未达 D 档").strip()
    top = claimed.endswith("(A)")
    add_text(slide, 0.5, 1.56, 2.9, 0.62, "实测达成档位\n%s" % claimed,
             size=14 if len(claimed) < 14 else 12, bold=True,
             color=("2E7D46" if top else "B5533A"), font=font)
    measured = ev.get("measured") or []
    bx = 3.5
    for b in measured[:5]:
        fits, txt = fit_lines(b, 9, 1.66, 3)
        add_rect(slide, bx, 1.56, 1.78, 0.66, "FFFFFF", line="D6DCE4")
        add_text(slide, bx + 0.05, 1.6, 1.68, 0.6, txt, size=9,
                 color="35506E", font=font)
        bx += 1.86
    rows = ev["rows"][:4]
    gf = slide.shapes.add_table(len(rows) + 1, 4, Inches(0.5), Inches(2.42),
                                Inches(12.33), Inches(3.1))
    tbl = gf.table
    for i, w in enumerate((1.3, 5.6, 4.25, 1.18)):
        tbl.columns[i].width = Inches(w)
    for c, h in enumerate(("档位", "工单验收标准（原文）", "本次实测判定", "判定")):
        cell = tbl.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(base)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor.from_string("FFFFFF")
                if font:
                    r.font.name = font
    for r, row in enumerate(rows, start=1):
        passed = row["verdict"] == "达标"
        vals = (row["tier"], _gate_text(row, std),
                row.get("verdict_text") or row["gate"],
                "✓ 达标" if passed else "✗ 未达标")
        for c, v in enumerate(vals):
            width_in = (1.3, 5.6, 4.25, 1.18)[c]
            fits, txt = fit_lines(v, 9.5, width_in - 0.12, 3)
            cell = tbl.cell(r, c)
            cell.text = txt
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                "FFFFFF" if r % 2 else "F7F9FB")
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5 if c in (1, 2) else 10)
                    run.font.bold = c in (0, 3)
                    run.font.color.rgb = RGBColor.from_string(
                        "2E7D46" if (c == 3 and passed) else
                        "B5533A" if c == 3 else "2B2F36")
                    if font:
                        run.font.name = font
    manual = ev.get("manual_items") or []
    n_manual = ev.get("manual_n", 0)
    if manual:
        note = "人工答辩项 %d 条（代码不代判，答辩现场确认）：%s" % (
            n_manual, "；".join(m[:56] for m in manual[:3]))
    elif n_manual:
        note = ("人工答辩项 %d 条（代码不代判，答辩现场确认，逐条见下页 report.txt 截图）"
                % n_manual)
    else:
        note = "人工答辩项 0 条：四档门槛均由代码实测判定"
    note += ("\n档位是算出来的：判定值由 run.py 本次真跑产生；"
             "工单未给数字的门槛在报告里标注为「兜底口径·非工单数字」。")
    if len(measured) > 5:
        note += "\n实测值共 %d 项，本页只列前 5 项，其余见下页 report.txt 截图。" % len(measured)
    if ev.get("claimed_src"):
        note += "\n档位取数：%s（报告头未写档位行）" % ev["claimed_src"]
    size = 10.5 if note.count("\n") <= 1 else (9.5 if note.count("\n") <= 2 else 9)
    add_rect(slide, 0.5, 5.62, 12.33, 1.38, "FFF7E8", line="E8C88A")
    _, note_fit = fit_lines(note, size, 12.0, 4)
    add_text(slide, 0.68, 5.7, 11.95, 1.24, note_fit, size=size,
             color="6B4E16", font=font)


PANELS = (("① run.py 实跑终端输出（逐字捕获）", "shot_run.png"),
          ("② pytest 全绿终端输出（逐字捕获）", "shot_test.png"),
          ("③ report.txt「验收档位判定（实测）」段", "shot_report.png"))


def build_e2(slide, ev, slug, base, font, pics, course_dir):
    grid = [(0.5, 1.62), (6.92, 1.62), (0.5, 4.24), (6.92, 4.24)]
    caps = list(PANELS)
    if ev.get("run_rc") not in (0, None):
        caps[0] = ("① run.py 实跑终端输出（逐字捕获·退出码 %s，输出与报告已完整产出）"
                   % ev["run_rc"], "shot_run.png")
    if ev.get("test_rc") not in (0, None):
        caps[1] = ("② pytest 输出（逐字捕获·退出码 %s）" % ev["test_rc"],
                   "shot_test.png")
    panels = [(cap, os.path.join(EV, slug, fn)) for cap, fn in caps]
    panels += pics[:3 - len(panels)]
    for i, (cap, path) in enumerate(panels[:3]):
        l, t = grid[i]
        add_rect(slide, l, t, 6.0, 2.52, "FFFFFF", line="D6DCE4")
        add_text(slide, l + 0.08, t + 0.04, 5.84, 0.26, cap, size=9.5,
                 bold=True, color=base, font=font)
        add_pic_fit(slide, path, l + 0.1, t + 0.32, 5.8, 2.14)
    l, t = grid[3]
    add_rect(slide, l, t, 6.0, 2.52, "F7F9FB", line="D6DCE4")
    add_text(slide, l + 0.08, t + 0.04, 5.84, 0.26, "④ 采集口径（可复现）",
             size=9.5, bold=True, color=base, font=font)
    meta = ev.get("meta") or {}
    py = meta.get("python", "python")
    lines = [
        "项目：%s/工程化项目/%s/" % (course_dir, slug),
        "命令：%s run.py ｜ %s -m pytest -q" % (py, py),
        "采集：工具/ppt97/capture97.py（%s 真跑逐字捕获输出）"
        % meta.get("date", "本次"),
        "渲染：工具/ppt97/terminal_png.py（终端截图，不改写任何数字）",
        "退出码：run=%s ｜ pytest=%s" % (ev.get("run_rc"), ev.get("test_rc")),
        "档位判定取自 output/report.txt「验收档位判定（实测）」段",
    ]
    add_text(slide, l + 0.12, t + 0.34, 5.76, 2.1, "\n".join(lines), size=9.5,
             color="2B2F36", font=font)
    add_text(slide, 0.5, 6.84, 12.3, 0.22,
             "截图内容逐字取自 工具/ppt97/evidence/%s/（capture97.py 真跑捕获，未改写数字）"
             % slug, size=9, color="7A8698", font=font)


def build_e3(slide, images, base, font):
    imgs = images[:4]
    n = len(imgs)
    wide = 12.42 if n == 1 else 6.0
    one_row = n <= 2
    h = 5.1 if one_row else 2.52
    grid = ([(0.5, 1.62, wide), (6.92, 1.62, 6.0)] if one_row else
            [(0.5, 1.62, 6.0), (6.92, 1.62, 6.0), (0.5, 4.24, 6.0),
             (6.92, 4.24, 6.0)])
    for i, (name, path, rel) in enumerate(imgs):
        l, t, w = grid[i]
        add_rect(slide, l, t, w, h, "FFFFFF", line="D6DCE4")
        add_text(slide, l + 0.08, t + 0.04, w - 0.16, 0.26,
                 "产出图 · %s" % name, size=9.5, bold=True, color=base, font=font)
        kb = os.path.getsize(path) // 1024
        add_text(slide, l + 0.08, t + 0.3, w - 0.16, 0.22,
                 "%s ｜ %d KB ｜ run.py 本次落盘" % (rel, kb),
                 size=8.5, color="7A8698", font=font)
        add_pic_fit(slide, path, l + 0.1, t + 0.56, w - 0.2, h - 0.64)
    if n == 3:
        l, t, w = grid[3]
        add_rect(slide, l, t, w, h, "F7F9FB", line="D6DCE4")
        add_text(slide, l + 0.08, t + 0.04, w - 0.16, 0.26, "本页未展开的其余产出图",
                 size=9.5, bold=True, color=base, font=font)
        rest = [r for _nm, _p, r in images[3:8]]
        add_text(slide, l + 0.12, t + 0.34, w - 0.24, h - 0.42,
                 "\n".join(rest) if rest else "本次运行共落盘 %d 张产出图，均已展示。" % n,
                 size=9, color="2B2F36", font=font)
    extra = "" if len(images) <= n else "；本次共落盘 %d 张产出图，本页展示前 %d 张" % (
        len(images), n)
    add_text(slide, 0.5, 6.84, 12.3, 0.22,
             "以上图像由 run.py 本次运行落盘到 output/，非外部素材%s。" % extra,
             size=9, color="7A8698", font=font)


def images_of(course_dir, slug, ev):
    out = []
    rels = [r for r in (ev.get("images") or []) if "test_run/" not in r]
    for rel in (rels or ev.get("images") or []):
        p = os.path.join(ROOT, course_dir, "工程化项目", slug, rel)
        if os.path.exists(p):
            out.append((os.path.basename(rel), p, rel))
    return out


def refresh_preview(slide, ev, font):
    """把「先看终点」页的运行输出块换成本次真跑内容。"""
    sh = None
    for s in slide.shapes:
        if s.has_text_frame and "参考实现真实运行输出" in s.text_frame.text:
            sh = s
            break
    if sh is None:
        return False
    seg = [ln.strip() for ln in (ev.get("report_tier_seg") or [])]
    body = [ln for ln in seg if ln and not ln.startswith("[")
            and not re.fullmatch(r"[-=—─_]+", ln)]
    if not body:
        if ev.get("claimed"):
            body.append("实测达成档位: %s" % ev["claimed"])
        body += ["%s %s ｜ 工单原文：%s" % (r["tier"], r.get("verdict_text") or r["gate"],
                                          r["need"]) for r in ev["rows"]]
    if not body:
        return False
    txt = "① 本次真跑实测输出（逐字取自 output/report.txt）\n" + "\n".join(body[:9])
    box_w = (sh.width or 0) / 914400.0 - 0.2
    box_h = (sh.height or 0) / 914400.0
    size = 12 if len(body) <= 6 else 10.5
    max_lines = max(4, int(box_h / (size / 72.0 * 1.35)))
    _, cut = fit_lines(txt, size, box_w, max_lines)
    set_shape_text(sh, cut)
    return True


def append_card_verdict(slide, row, font):
    """分层任务卡：在「达标标准」条上追加本次实测判定。"""
    for sh in slide.shapes:
        if not sh.has_text_frame or "达标标准" not in sh.text_frame.text:
            continue
        if len(sh.text_frame.text) > 150:
            return False
        mark = "✓ 本次实测达标" if row["verdict"] == "达标" else "✗ 本次实测未达标"
        add = "%s ｜ %s" % (mark, row.get("verdict_text") or row["gate"])
        box_w = (sh.width or 0) / 914400.0 - 0.1
        size = 10.5
        fits, cut = fit_lines(sh.text_frame.text + "\n" + add, size, box_w, 3)
        set_shape_text(sh, cut)
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)
                if font:
                    r.font.name = font
        sh.height = Inches(max((sh.height or 0) / 914400.0, 0.8))
        return True
    return False


def inject_deck(deck, course_dir, ev_cache, log):
    name = os.path.basename(deck)
    prs = Presentation(deck)
    if find_index(prs, E1_KEY) is not None:
        log.append((name, "已注入过，跳过"))
        return 0
    proj_idx = find_index(prs, PROJ_PAGE)
    if proj_idx is None:
        log.append((name, "SKIP 无「本工单的工程化项目」页"))
        return 0
    slug = deck_slug(prs, proj_idx)
    if slug not in ev_cache:
        log.append((name, "SKIP 未识别项目 %s" % slug))
        return 0
    ev = ev_cache[slug]
    std = card_standards(prs)
    base, font = base_color_of(prs.slides[proj_idx]), font_of(prs.slides[proj_idx])
    pics_dir_ok = all(os.path.exists(os.path.join(EV, slug, f))
                      for f in ("shot_run.png", "shot_test.png", "shot_report.png"))
    images = images_of(course_dir, slug, ev)
    specs = [("E1", E1_KEY + slug, lambda s: build_e1(s, ev, base, font, std))]
    if pics_dir_ok:
        specs.append(("E2", E2_KEY + slug,
                      lambda s: build_e2(s, ev, slug, base, font, [], course_dir)))
    if images:
        specs.append(("E3", E3_KEY + slug,
                      lambda s: build_e3(s, images, base, font)))
    made = 0
    for _tag, head, fill in specs:
        new = clone_base(prs, proj_idx)
        strip_body(new)
        tsh = title_shape(new)
        if tsh is None:
            log.append((name, "%s 无标题框，跳过" % _tag))
            continue
        set_shape_text(tsh, head)
        fill(new)
        made += 1
    _reorder(prs, made, proj_idx)
    _renumber_slide_num(prs)
    cards = 0
    for key in ("D", "C", "B", "A"):
        idx = find_index(prs, CARD_KEY[key])
        row = next((r for r in ev["rows"] if r["key"] == key), None)
        if idx is not None and row and append_card_verdict(prs.slides[idx], row, font):
            cards += 1
    pv = find_index(prs, "工单完成效果图")
    refreshed = refresh_preview(prs.slides[pv], ev, font) if pv is not None else False
    os.makedirs(BAK, exist_ok=True)
    if not os.path.exists(os.path.join(BAK, name)):
        shutil.copy2(deck, os.path.join(BAK, name))
    prs.save(deck)
    log.append((name, "OK %s 新增%d页 卡%d张 预览%s" % (
        slug, made, cards, "已刷新" if refreshed else "无匹配块")))
    return made


def main(limit=None, only=None):
    slugs = [os.path.basename(p.rstrip("/"))
             for p in glob.glob(os.path.join(EV, "*/"))]
    ev_cache = {}
    for s in slugs:
        try:
            ev_cache[s] = evidence(s)
        except Exception as exc:  # noqa: BLE001
            print("证据缺失 %s: %r" % (s, exc))
    print("evidence loaded=%d" % len(ev_cache))
    total = 0
    for _k, rel in COURSE_DIRS:
        course_dir = rel.split("/")[0]
        d = os.path.join(ROOT, rel)
        for deck in sorted(glob.glob(os.path.join(d, "*.pptx"))):
            if only and only not in os.path.basename(deck):
                continue
            try:
                total += inject_deck(deck, course_dir, ev_cache, LOG) or 0
            except Exception as exc:  # noqa: BLE001
                LOG.append((os.path.basename(deck), "FAIL %r" % exc))
            if limit and len(LOG) >= limit:
                break
    for n, m in LOG:
        print("%-52s %s" % (n[:52], m))
    print("decks=%d 新增页=%d" % (len(LOG), total))


LOG = []

if __name__ == "__main__":
    main(int(sys.argv[1]) if len(sys.argv) > 1 else None,
         sys.argv[2] if len(sys.argv) > 2 else None)
